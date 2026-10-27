import sys
from docx import Document
from docx.shared import Pt

path = sys.argv[1]
doc = Document(path)

print("=== SECTIONS ===")
for s in doc.sections:
    print("page", s.page_width, s.page_height, "margins L/R/T/B:",
          s.left_margin, s.right_margin, s.top_margin, s.bottom_margin,
          "orientation", s.orientation)

print("\n=== STYLES USED (in order of first appearance) ===")
seen = {}
for p in doc.paragraphs:
    sn = p.style.name if p.style else "None"
    key = (sn,)
    if key not in seen:
        fmt = []
        pf = p.paragraph_format
        fmt.append(f"align={pf.alignment}")
        fmt.append(f"spaceBefore={pf.space_before}")
        fmt.append(f"spaceAfter={pf.space_after}")
        fmt.append(f"lineSpacing={pf.line_spacing}")
        fmt.append(f"firstLine={pf.first_line_indent}")
        for r in p.runs[:3]:
            fmt.append(f"run(font={r.font.name},size={r.font.size},bold={r.bold},italic={r.italic})")
        seen[key] = (p.text[:70], " | ".join(fmt))
for k, v in seen.items():
    print(f"\nSTYLE: {k[0]}\n  text: {v[0]}\n  {v[1]}")

print("\n=== DEFAULT STYLE ===")
st = doc.styles['Normal']
print("font", st.font.name, st.font.size)
try:
    print("east asia", st.element.rPr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia'))
except Exception as e:
    print("no rFonts", e)
print("paragraph format:", st.paragraph_format.alignment, st.paragraph_format.line_spacing, st.paragraph_format.space_after)

print("\n=== TABLE STYLES ===")
for t in doc.tables:
    print("table style:", t.style.name if t.style else None, "rows", len(t.rows), "cols", len(t.columns))
