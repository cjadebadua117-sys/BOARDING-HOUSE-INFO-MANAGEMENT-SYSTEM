# -*- coding: utf-8 -*-
"""Write one DOCX per member with their module's DFD + ERD code.

Each document holds, for its module only:
  * the rendered DFD and ERD (so the code has a visible result beside it),
  * the exact Mermaid source of both diagrams, in a monospace box,
  * the one-line command that re-renders the PNG from the .mmd.

The .mmd files stay the source of truth; this just packages them for handing in.

Usage:
    venv\\Scripts\\python.exe tools\\build_module_diagram_docs.py
"""

import pathlib
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
DIAG = ROOT / "docs" / "final-project" / "diagrams"
DESKTOP = pathlib.Path.home() / "Desktop"

BODY = "Times New Roman"
MONO = "Consolas"
CODE_SIZE = 7.5
CODE_WIDTH = Inches(6.9)

SYSTEM = ("BHIMS: Boarding House Information Management System for Efficient "
          "Vacancy Search Among NLUC Bacnotan Students")
COURSE = "Business Information Systems & Systems Analysis and Design - BSIS 3B"

# letter, file owner tag, full name, module title
MEMBERS = [
    ("A", "Noemi", "Valdriz, Noemi G.",
     "Boarding House & Room Management with Verification"),
    ("B", "mac", "Layco, Mac Gabrielle L.",
     "Search & Inquiry"),
    ("C", "Jerome", "Bayubay, Jerome O.",
     "Booking, Acceptance & Ratings"),
    ("D", "cjade", "Badua, Christian Jade B.",
     "Accounts, Administration, Analytics & Reporting"),
]


# ---------------------------------------------------------------- docx helpers
def set_run_font(run, name, size, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if color:
        run.font.color.rgb = color


def para(doc, text="", align=None, bold=False, italic=False, size=12,
         before=0, after=4, name=BODY, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if text:
        set_run_font(p.add_run(text), name, size, bold, italic, color)
    return p


def cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table, color="D8C6B6", size=4):
    tbl_pr = table._tbl.tblPr
    xml = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        xml.append(node)
    tbl_pr.append(xml)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def code_box(doc, source):
    """One unbroken monospace box: a single table row that cannot split pages."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    no_split(table.rows[0])
    cell = table.cell(0, 0)
    cell.width = CODE_WIDTH
    shade(cell, "F7F3EC")
    cell_margins(cell)

    lines = source.rstrip("\n").splitlines()
    for index, line in enumerate(lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        # A literal empty run gives the paragraph no height on some renderers.
        set_run_font(p.add_run(line if line else " "), MONO, CODE_SIZE)
    return table


def figure(doc, png, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(png), width=CODE_WIDTH)
    para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=9, after=8)


# ---------------------------------------------------------------- build one doc
def build(letter, tag, name, title):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.8)
    s.top_margin = s.bottom_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    normal.element.rPr.rFonts.set(qn("w:cs"), BODY)
    normal.paragraph_format.space_after = Pt(4)

    cp = doc.core_properties
    cp.title = "BHIMS Module %s - %s diagram source" % (letter, title)
    cp.author = name
    cp.subject = COURSE
    cp.comments = "Module %s (%s) Mermaid source for the DFD and ERD" % (letter, tag)

    para(doc, SYSTEM, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True, after=2)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True, after=14)

    para(doc, "MODULE %s DIAGRAMS" % letter, align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=True, size=18, after=2)
    para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, after=2)
    para(doc, name, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=14)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    borders(table)
    rows = [
        ("Module", "%s - %s" % (letter, title)),
        ("Owner", name),
        ("Diagram files", "Module %s - %s - DFD (%s).mmd / .png\n"
                          "Module %s - %s - ERD (%s).mmd / .png"
                          % (letter, tag, title, letter, tag, title)),
        ("Rendered by", "Mermaid (flowchart for the DFD, erDiagram for the ERD)"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        for cell in row.cells:
            cell_margins(cell, 70, 70)
            shade(cell, "F7F3EC" if cell is row.cells[0] else "FFFFFF")
        set_run_font(row.cells[0].paragraphs[0].add_run(label), BODY, 10, bold=True)
        run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(run, BODY, 10)

    para(doc, "How to use this file", bold=True, size=13, before=16, after=3)
    para(doc, "1. The image above each code box is that diagram as rendered from the code "
              "below it - they are not hand drawings, so the two always agree.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
    para(doc, "2. Paste a code box into https://mermaid.live to edit it and export a new "
              "PNG or SVG without installing anything.", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
    para(doc, "3. To rebuild every diagram in the project from its .mmd file, run:",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)
    para(doc, r"venv\Scripts\python.exe tools\render_diagrams.py docs\final-project\diagrams",
         name=MONO, size=8.5, after=10)

    for kind, label, caption in (
        ("DFD", "Data Flow Diagram (DFD)",
         "Module %s Data Flow Diagram - the external entities, processes and data stores "
         "this module reads or writes." % letter),
        ("ERD", "Entity Relationship Diagram (ERD)",
         "Module %s Entity Relationship Diagram - only the tables this module writes, "
         "with the fields it manages." % letter),
    ):
        doc.add_page_break()
        para(doc, "%s - %s" % (kind, label), bold=True, size=14, before=0, after=6)
        stem = "Module %s - %s - %s (%s)" % (letter, tag, kind, title)
        png = DIAG / (stem + ".png")
        mmd = DIAG / (stem + ".mmd")
        if not png.exists() or not mmd.exists():
            raise SystemExit("missing diagram pair for %s" % stem)
        figure(doc, png, caption)
        para(doc, "Mermaid source (paste into mermaid.live)", bold=True, size=11,
             before=6, after=4)
        code_box(doc, mmd.read_text(encoding="utf-8"))

    out = DESKTOP / ("BHIMS Module %s - %s - Diagram Code.docx" % (letter, tag))
    doc.save(str(out))
    return out


def main():
    if not DESKTOP.is_dir():
        raise SystemExit("Desktop folder not found: %s" % DESKTOP)
    written = [build(*member) for member in MEMBERS]
    for path in written:
        print("%7.1f KB  %s" % (path.stat().st_size / 1024, path))


if __name__ == "__main__":
    sys.exit(main())
