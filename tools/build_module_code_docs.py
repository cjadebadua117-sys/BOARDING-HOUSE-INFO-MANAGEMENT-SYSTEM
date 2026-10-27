# -*- coding: utf-8 -*-
"""Write one DOCX per member holding the real Python code of their module.

The code is extracted from the running project by name (AST line ranges), not
retyped, so every line in the document is byte-identical to the file on disk:

    core/models.py   the module's tables and their helpers
    core/forms.py    the module's forms
    core/views.py    the module's request handlers and helpers
    core/urls.py     the routes that reach them

Templates are presentation, so they are listed by path (an index the member can
open) rather than printed as 1,000-line HTML pages.

Usage:
    venv\\Scripts\\python.exe tools\\build_module_code_docs.py
"""

import ast
import pathlib
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
DESKTOP = pathlib.Path.home() / "Desktop"

BODY = "Times New Roman"
MONO = "Consolas"
CODE_SIZE = 7.5
CODE_COLS = 108          # fits the 6.9" text column at 7.5 pt Consolas
CONTINUED = "    \u21b3 "  # marks a wrapped line; joining the pieces restores it
SHADE = "F7F3EC"
RULE = "D8C6B6"

SYSTEM = ("BHIMS: Boarding House Information Management System for Efficient "
          "Vacancy Search Among NLUC Bacnotan Students")
COURSE = "Business Information Systems & Systems Analysis and Design - BSIS 3B"

# ---------------------------------------------------------------------------
# The split. Every name below must exist in the file it is listed under.
# ---------------------------------------------------------------------------
MODULES = [
    {
        "letter": "A",
        "tag": "Noemi",
        "name": "Valdriz, Noemi G.",
        "title": "Boarding House & Room Management with Verification",
        "models": ["validate_image_size", "_upload_image_to", "_generated_image_path",
                   "validate_image_file", "BoardingHouse", "Room",
                   "annotate_room_bookings", "Amenity", "Barangay",
                   "barangay_display_map", "_reset_barangay_cache", "RoomPhoto",
                   "BACNOTAN_BARANGAYS", "BARANGAY_DISTANCES", "_barangay_display_cache"],
        "forms": ["BoardingHouseForm", "RoomForm"],
        "views": ["rooms_with_space", "visible_listings", "invalidate_listing_caches",
                  "landlord_dashboard", "landlord_boarding_house_list",
                  "boarding_house_create", "boarding_house_edit", "room_create",
                  "room_edit", "toggle_room_availability",
                  "landlord_boarding_house_detail", "admin_pending_listings",
                  "admin_approve_listing", "admin_reject_listing"],
        "urls": ["landlord_dashboard", "landlord_boarding_house_list",
                 "boarding_house_create", "boarding_house_edit",
                 "landlord_boarding_house_detail", "room_create", "room_edit",
                 "toggle_room_availability", "admin_pending_listings",
                 "admin_approve_listing", "admin_reject_listing"],
        "templates": ["landlord_dashboard.html", "landlord_boarding_house_list.html",
                      "landlord_boarding_house_detail.html", "boarding_house_form.html",
                      "room_form.html", "admin_pending_listings.html"],
    },
    {
        "letter": "B",
        "tag": "mac",
        "name": "Layco, Mac Gabrielle L.",
        "title": "Search & Inquiry",
        "models": ["Inquiry", "InquiryMessage", "MAX_MESSAGE_CHARS",
                   "LISTING_CACHE_KEY", "LANDING_CACHE_KEY"],
        "forms": ["ListingFilterForm", "InquiryForm"],
        "views": ["featured_listings", "home", "listing_list", "boarding_detail",
                  "listing_detail", "search_results", "send_inquiry", "inquiry_thread",
                  "inquiry_messages_json", "student_inquiry_list",
                  "landlord_inquiry_list", "student_dashboard", "chat_list",
                  "_with_latest_message", "_inquiry_has_new_message",
                  "_amenity_ids_from_names", "_amenity_ids_from",
                  "_matching_room_filter", "_apply_listing_filters",
                  "decorate_house_cards", "LISTINGS_PER_PAGE",
                  "LANDING_FEATURED_LIMIT", "LANDING_FEATURED_LANDLORDS"],
        "urls": ["home", "listing_list", "boarding_detail", "listing_detail",
                 "search_results", "send_inquiry",
                 "student_inquiry_thread", "landlord_inquiry_thread",
                 "inquiry_messages_json", "student_inquiry_list",
                 "landlord_inquiry_list", "student_dashboard", "chat_list"],
        "templates": ["landing.html", "boardings_list.html", "boardings_list_app.html",
                      "boarding_detail.html", "boarding_detail_app.html",
                      "search_results.html", "send_inquiry.html", "inquiry_thread.html",
                      "student_inquiry_list.html", "student_dashboard.html",
                      "landlord_inquiry_list.html", "includes/browse_content.html",
                      "includes/boarding_detail_content.html",
                      "includes/listing_card.html",
                      "includes/listing_grid_fragment.html",
                      "includes/search_grid_fragment.html"],
    },
    {
        "letter": "C",
        "tag": "Jerome",
        "name": "Bayubay, Jerome O.",
        "title": "Booking, Acceptance & Ratings",
        "models": ["Booking", "LandlordRating"],
        "forms": ["LandlordRatingForm"],
        "views": ["_rating_redirect", "rate_landlord", "accept_student",
                  "reject_student", "student_bookings", "landlord_bookings",
                  "admin_ratings", "landlord_my_ratings", "user_profile"],
        "urls": ["accept_student", "reject_student", "student_bookings",
                 "landlord_bookings", "landlord_my_ratings", "admin_ratings",
                 "rate_landlord", "user_profile"],
        "templates": ["student_bookings.html", "landlord_bookings.html",
                      "landlord_confirmation.html", "landlord_my_ratings.html",
                      "admin_ratings.html", "user_profile.html"],
    },
    {
        "letter": "D",
        "tag": "cjade",
        "name": "Badua, Christian Jade B.",
        "title": "Accounts, Administration, Analytics & Reporting",
        "models": ["UserManager", "User", "AdminActionLog", "log_admin_action",
                   "SiteSetting", "site_setting", "UserReport",
                   "_invalidate_listing_caches"],
        "forms": ["sync_display_name", "StudentRegistrationForm",
                  "LandlordRegistrationForm", "StudentProfileEditForm",
                  "BHIMSPasswordChangeForm", "LandlordCreateForm",
                  "LandlordEditForm", "StudentProfileUpdateForm",
                  "LandlordProfileUpdateForm", "LandlordProfilePictureForm",
                  "BarangayForm", "SystemSettingsForm", "UserReportForm"],
        "views": ["admin_required", "landlord_required", "student_required",
                  "_dashboard_for", "_redirect_if_authenticated",
                  "_safe_redirect_target", "_client_ip", "_cache_get", "_cache_set",
                  "_hit", "_throttle_login", "_throttle_key", "_throttle_generic",
                  "_lockout_state", "_record_login_failure", "_clear_login_failures",
                  "handler404", "handler500", "student_register", "accounts_login",
                  "landlord_register", "student_login", "landlord_login",
                  "admin_login", "logout_view", "profile_view", "_serialize_profile",
                  "update_profile", "upload_profile_picture", "change_password",
                  "report_user", "profile", "edit_profile", "update_avatar",
                  "admin_dashboard", "admin_analytics", "admin_create_landlord",
                  "admin_reports", "admin_report_detail", "admin_report_update",
                  "admin_landlord_list", "admin_landlord_detail",
                  "admin_edit_landlord", "toggle_landlord_active",
                  "admin_delete_landlord", "admin_student_list",
                  "admin_delete_student", "admin_barangay_list",
                  "admin_barangay_edit", "admin_barangay_delete", "admin_settings",
                  "health_check", "LOGIN_FAIL_LIMIT", "LOGIN_WINDOW", "LOGIN_LOCK",
                  "_ROLE_DASHBOARDS"],
        "urls": ["student_register", "student_login", "landlord_register",
                 "landlord_login", "landlord_login_redirect", "admin_login",
                 "logout", "login", "profile", "update_profile", "edit_profile",
                 "update_avatar", "change_password", "report_user",
                 "admin_dashboard", "admin_analytics", "admin_create_landlord",
                 "admin_landlord_list", "admin_landlord_detail",
                 "admin_edit_landlord", "toggle_landlord_active",
                 "admin_delete_landlord", "admin_student_list",
                 "admin_delete_student", "admin_barangay_list",
                 "admin_barangay_edit", "admin_barangay_delete", "admin_settings",
                 "admin_reports", "admin_report_detail", "admin_report_update",
                 "health_check"],
        "templates": ["auth/base.html", "auth/login_student.html",
                      "auth/login_landlord.html", "auth/login_admin.html",
                      "auth/_field.html", "auth/_eye_toggle.html",
                      "auth/_split_art.html", "includes/login_form_footer.html",
                      "admin_dashboard.html", "admin_analytics.html",
                      "admin_reports.html", "admin_report_detail.html",
                      "admin_landlord_list.html", "admin_landlord_detail.html",
                      "admin_student_list.html", "admin_barangays.html",
                      "admin_settings.html", "admin_honeypot.html", "profile.html",
                      "edit_profile.html", "change_password.html",
                      "report_user.html", "password_reset.html",
                      "password_reset_done.html", "password_reset_confirm.html",
                      "password_reset_complete.html", "404.html", "500.html"],
    },
]


# ---------------------------------------------------------------------------
# Source extraction
# ---------------------------------------------------------------------------
def top_level_blocks(path):
    """name -> (start, end) line numbers for every top-level def/class/assign."""
    source = path.read_text(encoding="utf-8")
    tree = ast.parse(source)
    blocks = {}
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            start = min([node.lineno] + [d.lineno for d in node.decorator_list])
            kind = "class" if isinstance(node, ast.ClassDef) else "def"
            blocks[node.name] = (start, node.end_lineno, kind)
        elif isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name):
                    blocks.setdefault(target.id, (node.lineno, node.end_lineno, "constant"))
                    break
        elif isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name):
            blocks.setdefault(node.target.id, (node.lineno, node.end_lineno, "constant"))
    return blocks


def normalise(text):
    return text.replace("\t", "    ")


def module_parts(path, names):
    """Header plus the requested blocks, in file order, as (caption, text)."""
    lines = normalise(path.read_text(encoding="utf-8")).splitlines()
    blocks = top_level_blocks(path)
    missing = [n for n in names if n not in blocks]
    if missing:
        raise SystemExit("%s is missing: %s" % (path.name, ", ".join(missing)))

    resolved = sorted({blocks[n] for n in names})
    labels = {}
    for name in names:
        labels.setdefault(blocks[name], []).append(name)

    tree = ast.parse("\n".join(lines))
    first_def = min(node.lineno for node in tree.body
                    if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)))
    parts = [("Imports and module constants", "\n".join(lines[:first_def - 1]).rstrip())]
    for start, end, kind in resolved:
        if start < first_def:
            continue  # already inside the header
        text = "\n".join(lines[start - 1:end]).rstrip()
        parts.append(("%s %s  (%s:%d)" % (kind, ", ".join(sorted(labels[(start, end, kind)])),
                                          path.name, start), text))
    return parts


def url_entries(names):
    """(pattern, view, name, source) for the requested urlpatterns entries."""
    path = ROOT / "core" / "urls.py"
    lines = normalise(path.read_text(encoding="utf-8")).splitlines()
    tree = ast.parse("\n".join(lines))
    found = {}
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Name):
            continue
        if node.func.id != "path":
            continue
        kwargs = {kw.arg: kw.value for kw in node.keywords if kw.arg}
        label = kwargs.get("name")
        if not isinstance(label, ast.Constant):
            continue
        if label.value not in names:
            continue
        pattern = node.args[0].value if node.args and isinstance(node.args[0], ast.Constant) else ""
        view = ast.unparse(node.args[1]) if len(node.args) > 1 else ""
        view = view.replace("views.", "")
        start = node.lineno
        # keep a comment block that sits directly above the route
        while start > 1 and lines[start - 2].lstrip().startswith("#"):
            start -= 1
        found[label.value] = (pattern, view, label.value,
                              "\n".join(lines[start - 1:node.end_lineno]).rstrip())
    missing = [n for n in names if n not in found]
    if missing:
        raise SystemExit("core/urls.py is missing: %s" % ", ".join(missing))
    return [found[n] for n in names if n in found]


# ---------------------------------------------------------------------------
# DOCX helpers
# ---------------------------------------------------------------------------
def set_font(run, name, size, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if color is not None:
        run.font.color.rgb = color


def para(doc, text="", align=None, bold=False, italic=False, size=12,
         before=0, after=4, name=BODY, color=None, indent=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.left_indent = Inches(indent)
    if text:
        set_font(p.add_run(text), name, size, bold, italic, color)
    return p


def shade_para(p, fill):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), RULE)
    border.append(left)
    ppr.append(border)


def display_lines(text):
    """Source lines, each split so nothing runs past the right margin."""
    out = []
    for line in normalise(text).splitlines() or [""]:
        line = line.rstrip()
        if len(line) <= CODE_COLS:
            out.append(line or " ")
            continue
        out.append(line[:CODE_COLS])
        rest = line[CODE_COLS:]
        width = CODE_COLS - len(CONTINUED)
        while rest:
            out.append(CONTINUED + rest[:width])
            rest = rest[width:]
    return out


def code_block(doc, text, size=CODE_SIZE):
    lines = display_lines(text)
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Inches(0.08)
        shade_para(p, SHADE)
        set_font(p.add_run(line), MONO, size)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), RULE)
        borders.append(node)
    tbl_pr.append(borders)

    for cell, text in zip(t.rows[0].cells, header):
        cell.width = Inches(widths[header.index(text)])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        set_font(p.add_run(text), BODY, 9, bold=True)
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "EFE6DC")
        tc_pr.append(shd)
    for row in rows:
        cells = t.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            set_font(p.add_run(value), BODY, 9)
    return t


def add_page_number_footer(doc, label):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(label + "  -  page "), BODY, 8, color=RGBColor(0x66, 0x55, 0x48))
    run = p.add_run()
    set_font(run, BODY, 8, color=RGBColor(0x66, 0x55, 0x48))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ---------------------------------------------------------------------------
# Build one document
# ---------------------------------------------------------------------------
def build(module):
    models_parts = module_parts(ROOT / "core" / "models.py", module["models"])
    forms_parts = module_parts(ROOT / "core" / "forms.py", module["forms"])
    views_parts = module_parts(ROOT / "core" / "views.py", module["views"])
    routes = url_entries(module["urls"])

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    normal.element.rPr.rFonts.set(qn("w:cs"), BODY)
    normal.paragraph_format.space_after = Pt(4)

    cp = doc.core_properties
    cp.title = "BHIMS Module %s - %s - source code" % (module["letter"], module["title"])
    cp.author = module["name"]
    cp.subject = COURSE
    cp.comments = "The Python code behind Module %s (%s)" % (module["letter"], module["tag"])

    add_page_number_footer(doc, "BHIMS - Module %s (%s)" % (module["letter"], module["tag"]))

    para(doc, SYSTEM, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True, after=2)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True, after=18)

    para(doc, "MODULE %s" % module["letter"], align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=True, size=20, after=2)
    para(doc, module["title"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, after=4)
    para(doc, "Source code", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12, after=16)
    para(doc, "Presented by", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, after=1)
    para(doc, module["name"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, after=26)

    total_lines = sum(len(text.splitlines()) for _, text in models_parts + forms_parts + views_parts)
    total_lines += sum(len(src.splitlines()) for _, _, _, src in routes)

    para(doc, "What this document contains", bold=True, size=13, after=4)
    para(doc, "This is the Python code that implements Module %s in the running system. "
              "Every block was copied out of the project file by its definition name, not "
              "retyped, so the line numbers in each heading still point at the same code in "
              "the repository. Lines wider than the page are continued on the next line with "
              "a \u21b3 marker, and joining the pieces gives the line back exactly. Templates are "
              "listed at the end for reference."
              % module["letter"], align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)

    table(doc,
          ["File", "What is printed here", "Blocks"],
          [["core/models.py", "The module's database tables, fields and model helpers",
            str(len(models_parts))],
           ["core/forms.py", "The forms the module validates input with", str(len(forms_parts))],
           ["core/views.py", "The request handlers and helpers that carry the feature",
            str(len(views_parts))],
           ["core/urls.py", "The routes that reach those handlers", str(len(routes))]],
          widths=[1.5, 4.4, 0.8])
    para(doc, "Total: %d lines of Python in this document." % total_lines,
         italic=True, size=10, before=6, after=0)

    # ---------------------------------------------------------------- models
    page_break(doc)
    para(doc, "1. Models", bold=True, size=16, after=2)
    para(doc, "core/models.py - the tables this module owns, plus the helpers that keep them "
              "consistent.", italic=True, size=10, after=6)
    for caption, text in models_parts:
        para(doc, caption, italic=True, size=8.5, before=6, after=2, name=MONO,
             color=RGBColor(0x6b, 0x4a, 0x32))
        code_block(doc, text)

    # ---------------------------------------------------------------- forms
    page_break(doc)
    para(doc, "2. Forms", bold=True, size=16, after=2)
    para(doc, "core/forms.py - server-side validation for everything this module accepts.",
         italic=True, size=10, after=6)
    for caption, text in forms_parts:
        para(doc, caption, italic=True, size=8.5, before=6, after=2, name=MONO,
             color=RGBColor(0x6b, 0x4a, 0x32))
        code_block(doc, text)

    # ---------------------------------------------------------------- views
    page_break(doc)
    para(doc, "3. Views", bold=True, size=16, after=2)
    para(doc, "core/views.py - the handlers behind this module's pages and actions, in the "
              "order they appear in the file.", italic=True, size=10, after=6)
    for caption, text in views_parts:
        para(doc, caption, italic=True, size=8.5, before=6, after=2, name=MONO,
             color=RGBColor(0x6b, 0x4a, 0x32))
        code_block(doc, text)

    # ---------------------------------------------------------------- urls
    page_break(doc)
    para(doc, "4. URL routes", bold=True, size=16, after=2)
    para(doc, "core/urls.py - each route this module serves, with the view it calls.",
         italic=True, size=10, after=6)
    table(doc, ["URL pattern", "View", "Route name"],
          [["/".join(["", pattern]) if not pattern.startswith("/") else pattern, view, name]
           for pattern, view, name, _ in routes],
          widths=[2.3, 2.6, 1.8])
    para(doc, "", after=4)
    code_block(doc, "urlpatterns = [\n" + "\n".join(
        src for _, _, _, src in routes) + "\n]")

    # ---------------------------------------------------------------- files
    page_break(doc)
    para(doc, "5. Templates and assets this module renders", bold=True, size=16, after=4)
    para(doc, "The HTML is presentation, so it is listed rather than printed. Open these under "
              "core/templates/core/ to see the pages that display this module's work.",
         italic=True, size=10, after=6)
    table(doc, ["Template", "Lines"],
          [[name, str(_line_count(ROOT / "core" / "templates" / "core" / name))]
           for name in module["templates"]],
          widths=[4.6, 0.8])
    para(doc, "Shared layout used by every module: base.html, app_base.html, "
              "includes/navbar.html, includes/footer.html.", italic=True, size=9, before=8)

    out = DESKTOP / ("BHIMS Module %s - %s - Feature Code.docx" % (module["letter"], module["tag"]))
    doc.save(str(out))
    return out, total_lines, len(models_parts), len(forms_parts), len(views_parts), len(routes)


def _line_count(path):
    try:
        return len(normalise(path.read_text(encoding="utf-8")).splitlines())
    except OSError:
        return -1


def main():
    if not DESKTOP.is_dir():
        raise SystemExit("Desktop folder not found: %s" % DESKTOP)
    for module in MODULES:
        out, lines, m, f, v, u = build(module)
        print("%7.1f KB  %-46s %4d code lines (models %d, forms %d, views %d, urls %d)"
              % (out.stat().st_size / 1024, out.name, lines, m, f, v, u))


if __name__ == "__main__":
    sys.exit(main())
