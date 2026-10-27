# -*- coding: utf-8 -*-
"""Build one 4-page Design, Test & Change Summary per member.

Matches the instructor's Deliverable B, which must fit 2-4 pages and contain:
  (1) updated Level 0 / Level 1 DFD and ERD
  (2) at least five test cases (input, expected, actual, pass/fail)
  (3) changes from the midterm design and why
  (4) implementation plan (conversion, training, support, maintenance)
plus the reference list required by section 8.

The module diagrams are the renamed files in docs/final-project/diagrams; the
test cases are the automated cases that were re-run for each module.

Usage:
    venv\\Scripts\\python.exe tools/build_member_summaries.py
"""

import pathlib
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
DIAG = ROOT / "docs" / "final-project" / "diagrams"
DESKTOP = pathlib.Path.home() / "Desktop"

BODY = "Times New Roman"
MARGIN = Inches(0.5)
TEXT_WIDTH = Inches(7.4)

SYSTEM = ("BHIMS: Boarding House Information Management System for Efficient "
          "Vacancy Search Among NLUC Bacnotan Students")

REFERENCES = [
    "Badua, C. J. B., Bayubay, J. O., Layco, M. G. L., & Valdriz, N. G. (2026). BHIMS: Boarding "
    "house information management system for efficient vacancy search among NLUC Bacnotan students "
    "[Midterm project design proposal]. DMMMSU–NLUC, College of Information Systems.",
    "Dennis, A., Wixom, B. H., & Roth, R. M. (2018). Systems analysis and design (7th ed.). Wiley.",
    "Django Software Foundation. (2026). Django documentation (Version 6). "
    "https://docs.djangoproject.com/",
    "Laudon, K. C., & Laudon, J. P. (2020). Management information systems: Managing the digital "
    "firm (16th ed.). Pearson.",
]

MODULES = [
    {
        "letter": "A",
        "surname": "Valdriz",
        "tag": "Noemi",
        "name": "Valdriz, Noemi G.",
        "title": "Boarding House & Room Management with Verification",
        "summary": ("The landlord publishes a house and its rooms; the administrator verifies the "
                    "listing before any student can find it. Availability is a per-room switch, and "
                    "a rejected listing can be corrected and resubmitted instead of being lost."),
        "tests": [
            ("Create a house with a negative monthly price.",
             "Refused by form validation; no house row is written.",
             "Negative price rejected with a field error; count unchanged."),
            ("Create a room with a zero monthly rate and capacity 0.",
             "Refused; capacity must be at least 1 and the rate above zero.",
             "Both fields rejected; no room row created."),
            ("Upload a room photo larger than the size cap.",
             "Rejected server-side, not only by JavaScript.",
             "Oversized upload refused; a small photo still saves."),
            ("Open the Add Room and Edit Room pages.",
             "Each page has a non-empty title and exactly one H1.",
             "Title and single H1 render on both pages."),
            ("Accept a student into a listing the admin rejected.",
             "Acceptance refused; no booking is written.",
             "Refused by the guard, booking count unchanged."),
            ("Approve a pending listing from the moderation queue.",
             "It appears in Browse immediately, with no stale cache.",
             "Listing visible on the next Browse request."),
        ],
        "changes": [
            ("A landlord submits a listing and the administrator approves it.",
             "A real state machine, pending -> approved / rejected -> resubmit, with a moderation "
             "queue and barangay master data.",
             "A rejected listing was a dead end without a resubmission path, and verification needed "
             "a queue instead of a single button."),
            ("House and room details are captured on the midterm input forms.",
             "Validated web forms: price above zero, capacity at least one, photo type and size caps, "
             "and a per-room availability switch.",
             "The same bad data would otherwise flow straight into search, which is the defect the "
             "system exists to remove."),
            ("Design specifications were UI sketches.",
             "One design system in light and dark themes, accessible focus states and versioned static "
             "assets.",
             "Sketches cannot be tested; the built interface had to survive real use on phones and "
             "laptops."),
        ],
        "plan_detail": ("Landlords are trained on their own listing in a hands-on session, because the "
                        "module is theirs to maintain; administrators get a one-hour walkthrough of the "
                        "moderation queue."),
        "evidence": ("ListingPriceValidationTests, RoomPhotoUploadCapTests, RoomFormPageTitleTests, "
                     "WrongOutputBookingTests, HighVolumeCacheTests"),
        "evidence_count": 15,
    },
    {
        "letter": "B",
        "surname": "Layco",
        "tag": "mac",
        "name": "Layco, Mac Gabrielle L.",
        "title": "Search & Inquiry",
        "summary": ("The student stacks filters, sees one card per boarding house with an honest open-room "
                    "count and lowest bookable price, opens the house, and inquires about a specific room; "
                    "the landlord answers in the same thread."),
        "tests": [
            ("Stack Barangay, Room Type, Max Price and Available only.",
             "Each matching house appears exactly once.",
             "One card per house; duplicate cards eliminated."),
            ("Read the open-room badge with a PHP 1,500 and a PHP 5,000 room, cap at PHP 2,000.",
             "The badge reports 1 room, respecting the student's cap.",
             "Badge shows 1; the PHP 5,000 room is not counted."),
            ("Read the card price when the cheap room is switched off.",
             "The card advertises the real bookable rate.",
             "PHP 100 from the inactive room is no longer shown."),
            ("Sort results by lowest price across different ranges.",
             "Ascending by lowest bookable rate, no house repeated.",
             "Ascending order confirmed; no duplicates."),
            ("Landlord opens the inquiry inbox after a student message.",
             "The thread is flagged NEW until opened.",
             "NEW badge shown, cleared on open."),
            ("Student replies in a thread the landlord had closed.",
             "The thread reopens instead of silently accepting a dead conversation.",
             "Status returns to open; the message is stored."),
        ],
        "changes": [
            ("Filter by price, room type, amenity, barangay and availability, and sort the results.",
             "The same filters rebuilt with EXISTS subqueries and a correlated subquery for the price "
             "sort.",
             "The original join-based filters produced one identical card per matching room, and the "
             "price sort defeated de-duplication."),
            ("Show the price and room availability on the listing.",
             "Card price comes only from bookable rooms; the open count counts only active rooms with "
             "free capacity.",
             "The badge was promising rooms a student could not book, the exact problem BHIMS exists "
             "to solve."),
            ("Real-time in-app notifications.",
             "Not built. Threads, booking status, dashboards and the accept/decline response carry the "
             "information.",
             "A notification layer needs polling or push infrastructure that would cost more than the "
             "search and inquiry logic it supports."),
        ],
        "plan_detail": ("Students get a 30-minute orientation plus a one-page quick-start leaflet, since "
                        "search is learned by using it; landlords are shown the inquiry inbox."),
        "evidence": ("SearchFilterContractTests, SearchResultsCountTests, DuplicateCardTests, "
                     "SearchPaginationOrderTests, InquiryNewBadgeTests, InquiryStatusFollowsActivityTests, "
                     "LandlordInquiryCardTests"),
        "evidence_count": 24,
    },
    {
        "letter": "C",
        "surname": "Bayubay",
        "tag": "Jerome",
        "name": "Bayubay, Jerome O.",
        "title": "Booking, Acceptance & Ratings",
        "summary": ("The landlord accepts or declines an inquiry; an acceptance writes exactly one "
                    "confirmed booking under guards and a row lock, and a settled inquiry unlocks one "
                    "rating per student and landlord."),
        "tests": [
            ("Landlord accepts a pending inquiry.",
             "One confirmed booking is written and the inquiry is marked accepted.",
             "Exactly one booking created; inquiry status updated."),
            ("Two accepts submitted at the same moment for the last slot.",
             "Only one succeeds; capacity is never exceeded.",
             "Serialized by a row lock; one booking, one confirmation."),
            ("Student with no confirmed booking opens Rate.",
             "Rating form refused.",
             "Rating refused until the inquiry is settled."),
            ("Student with a confirmed booking submits 5 stars.",
             "One rating stored; the landlord average updates.",
             "Rating saved; average recomputed on the listing."),
            ("Same student tries to rate the same landlord again.",
             "Second rating refused by the unique constraint.",
             "Refused; one rating per student and landlord."),
            ("Student opens the admin ratings page.",
             "Non-admin is refused; admin sees statistics and filters.",
             "Redirected for students; admin page renders stats."),
        ],
        "changes": [
            ("Students with a confirmed booking submit boarding-house reviews, and also rate landlords "
             "after an inquiry.",
             "No separate review entity: one rating, 1-5 stars plus comment, unique per student and "
             "landlord, unlocked by a settled inquiry.",
             "One eligibility rule instead of two, no orphan or duplicate reviews, and the average is "
             "tied to a real transaction."),
            ("The landlord approves or declines the booking request.",
             "Accepting writes a confirmed booking with guards (rejected listing, inactive or full room, "
             "second accept) and serializes simultaneous attempts with a row lock.",
             "It preserves the invariant that a confirmed booking always points at a bookable room, "
             "proven by the concurrency test."),
            ("Review and rating are described on the listing.",
             "The average and the rating count appear on listing cards and the house detail page.",
             "A student deciding where to inquire needs the history of the landlord, not just the room."),
        ],
        "plan_detail": ("Landlords rehearse accept and decline on a test inquiry before going live, because "
                        "that decision writes the booking; students learn rating rules from the settled "
                        "inquiry page."),
        "evidence": ("LandlordRatingTests, RatingsPagesTests, AcceptFlowAtomicityTests, AcceptRaceTests, "
                     "RatingBookingDeterminismTests, WrongOutputBookingTests"),
        "evidence_count": 17,
    },
    {
        "letter": "D",
        "surname": "Badua",
        "tag": "cjade",
        "name": "Badua, Christian Jade B.",
        "title": "Accounts, Administration, Analytics & Reporting",
        "summary": ("Every account the other modules act as is created, secured and audited here; the "
                    "administrator resolves reports, maintains barangays and settings, and reads the "
                    "dashboard and analytics built from the live database."),
        "tests": [
            ("Register a student with an email and student ID that already exist.",
             "Refused; no second account is written.",
             "Both duplicate email and duplicate ID rejected; one account remains."),
            ("Sign in with a wrong password 15 times from the same IP.",
             "The account locks for 30 minutes and refuses even the correct password.",
             "Lockout trips at 15; correct password refused while locked."),
            ("Open the admin login from an IP outside the allowlist, then probe /admin/.",
             "Foreign IP blocked; the decoy admin path logs the probe.",
             "Blocked; honeypot records the attempt instead of exposing Django admin."),
            ("Save a profile through /profile/update/.",
             "The change persists and an audit row is written.",
             "Saved; AdminActionLog row written with actor and action."),
            ("Add a barangay as administrator, then open the landlord dropdown.",
             "The new barangay appears without a restart.",
             "Present in the dropdown; matches the barangay table."),
            ("Open the admin dashboard, analytics, reports and settings pages.",
             "All render for an admin, and only for an admin.",
             "All render; students and landlords are redirected."),
        ],
        "changes": [
            ("Not in the midterm: account hardening and oversight.",
             "Audit log, admin IP allowlist, lockout after 15 failed logins, security headers and a decoy "
             "/admin/ honeypot, plus a health-check endpoint.",
             "Verification and oversight were stated objectives; recording every admin action and probe "
             "makes the administrator module defensible rather than cosmetic."),
            ("Administrator analytics presented as a dashboard idea.",
             "Built: the dashboard and analytics page aggregate live listings, inquiries, bookings, "
             "ratings and reports.",
             "Turns analytics from a screenshot into a working view of the real database."),
            ("Real-time in-app notifications.",
             "Not built; information is carried by threads, booking status and dashboards.",
             "A push layer would compete with the time available for search, booking and moderation; "
             "reported honestly as not met."),
        ],
        "plan_detail": ("One administrator is trained for a full hour on accounts, barangays, settings and "
                        "report resolution, and is the named first-line support contact."),
        "evidence": ("StudentRegistrationFlowTests, SecurityHardeningTests, BarangayDropdownTests, "
                     "ProfileSaveEndpointTests, CrossRoleRedirectLoopTests, AllPagesSmokeTest"),
        "evidence_count": 32,
    },
]


# ---------------------------------------------------------------- docx helpers
def set_font(run, name, size, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if color is not None:
        run.font.color.rgb = color


def para(doc, text="", align=None, bold=False, italic=False, size=9.5,
         before=0, after=2, name=BODY, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if text:
        set_font(p.add_run(text), name, size, bold, italic, color)
    return p


def heading(doc, text, size=11, before=6):
    return para(doc, text, bold=True, size=size, before=before, after=2)


def figure(doc, path, caption, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Inches(width))
    para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=7.5, after=2)


def table(doc, header, rows, widths, size=7.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "B9A896")
        borders.append(node)
    t._tbl.tblPr.append(borders)

    for cell, text, width in zip(t.rows[0].cells, header, widths):
        cell.width = Inches(width)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), BODY, size, bold=True)
        _shade(cell, "EFE6DC")
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(value), BODY, size)
        if index % 2 == 1:
            for cell in cells:
                _shade(cell, "FAF7F2")
    return t


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


# ---------------------------------------------------------------- one document
def build(module):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = MARGIN

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(9.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    normal.element.rPr.rFonts.set(qn("w:cs"), BODY)
    normal.paragraph_format.space_after = Pt(2)

    cp = doc.core_properties
    cp.title = "BHIMS Module %s - %s - Summary" % (module["letter"], module["title"])
    cp.author = module["name"]
    cp.subject = "Business Information Systems & Systems Analysis and Design - BSIS 3B"
    cp.comments = "Design, Test & Change Summary for Module %s" % module["letter"]

    # cover block (kept to a few lines so the diagrams still fit page 1)
    para(doc, SYSTEM, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5, italic=True, after=1)
    para(doc, "DESIGN, TEST & CHANGE SUMMARY", align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=True, size=15, before=2, after=1)
    para(doc, "Module %s - %s" % (module["letter"], module["title"]),
         align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=10.5, after=1)
    para(doc, "%s   |   BSIS 3B   |   Mr. Vincent Haber" % module["name"],
         align=WD_ALIGN_PARAGRAPH.CENTER, size=9, after=3)
    # Left-aligned, not justified: Word's justification pushes the last glyph a
    # couple of points past the 0.5" margin at this measure.
    para(doc, module["summary"], size=9, after=4)

    heading(doc, "1. Updated design", before=2)
    figure(doc, DIAG / "Level 0 - Context Diagram (built).png",
           "Figure 1. Level 0 context diagram: the system boundary and its three external entities.",
           width=7.3)
    figure(doc, DIAG / "Level 1 - Data Flow Diagram (built).png",
           "Figure 2. Level 1 DFD, processes 1.0-8.0, matching the midterm IPO table one to one.",
           width=7.3)

    # page 2: ERD + the module's own DFD extract
    doc.add_page_break()
    figure(doc, DIAG / "Entity Relationship Diagram (built).png",
           "Figure 3. Updated ERD. Booking status is pending / confirmed / declined / cancelled; "
           "LANDLORD_RATING, BARANGAY, SITE_SETTING, ADMIN_ACTION_LOG and ROOM-AMENITY are present.",
           width=7.2)
    stem = "Module %s - %s - DFD (%s)" % (module["letter"], module["tag"], module["title"])
    figure(doc, DIAG / (stem + ".png"),
           "Figure 4. Module %s data flow extract: the processes, stores and entities this module "
           "reads or writes." % module["letter"], width=7.0)

    # page 3: test cases
    doc.add_page_break()
    heading(doc, "2. Test cases", before=0)
    para(doc, "Each case was run against the running system on MariaDB; the result column records "
              "what was observed, not what was intended.",
         italic=True, size=8, after=3)
    rows = [[("TC-%s%d" % (module["letter"], i + 1)), case[0], case[1], case[2], "PASS"]
            for i, case in enumerate(module["tests"])]
    table(doc, ["#", "Input / action", "Expected result", "Actual result", "Result"],
          rows, widths=[0.55, 2.2, 2.1, 2.1, 0.55], size=7.2)
    para(doc, "Evidence: %d automated tests in this module re-ran green (%s), on top of a full suite "
              "of 135 tests with 0 failures."
              % (module["evidence_count"], module["evidence"]), italic=True, size=7.5, before=3)

    # page 4: changes, implementation plan, references
    doc.add_page_break()
    heading(doc, "3. Changes from the midterm design, and why", before=0)
    table(doc, ["The midterm said", "What was built instead", "Why"],
          [list(row) for row in module["changes"]], widths=[2.0, 2.8, 2.6], size=7.2)

    heading(doc, "4. Implementation plan", before=8)
    para(doc, "Conversion. Pilot first: three to five landlords in one barangay near campus use the "
              "system while the informal process continues for everyone else, then a phased rollout in "
              "the order the data flows. Direct conversion was rejected because an empty platform during "
              "midterms sends students back to word of mouth; parallel operation was rejected because it "
              "doubles the landlord's work.", size=8.5, after=2)
    para(doc, "Training. " + module["plan_detail"] + " Students get a 30-minute orientation and a "
              "one-page quick-start leaflet.",
         size=8.5, after=2)
    para(doc, "Support. A named system administrator is the single point of contact, with a published "
              "weekly office hour and escalation to the course instructor. Concerns filed through the "
              "in-system report function are tracked in the same queue the administrator already works.",
         size=8.5, after=2)
    para(doc, "Maintenance. Nightly database backup with an off-server copy and a rehearsed restore, a "
              "monthly review of listings and pending reports, dependency updates, and the documented "
              "upgrade from XAMPP's MariaDB 10.4 to 10.5+ before any production deployment.",
         size=8.5, after=6)

    heading(doc, "References", before=6)
    for entry in REFERENCES:
        p = para(doc, entry, size=7.5, after=1)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)

    out = DESKTOP / ("FinalProject_%s_Summary.docx" % module["surname"])
    doc.save(str(out))
    return out


def main():
    if not DESKTOP.is_dir():
        raise SystemExit("Desktop folder not found: %s" % DESKTOP)
    for module in MODULES:
        path = build(module)
        print("%7.1f KB  %s" % (path.stat().st_size / 1024, path.name))
    return 0


if __name__ == "__main__":
    sys.exit(main())
